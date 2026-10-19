#!/usr/bin/env python3
"""Combine the Process & Data Dictionary workbook and the Data Library workbook
into one file: dictionary tabs first, then the library tabs.

The library's "Read Me" is renamed "Library Read Me"; its "Cleanup Backlog" is
dropped and its Note column is appended to the dictionary's Cleanup Backlog tab
(same 21 items, same order). The Contents tab gains rows for the library tabs.

Usage: combine_workbooks.py <dictionary.xlsx> <library.xlsx> <out.xlsx>
"""
import sys
from copy import copy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

DICT, LIB, OUT = sys.argv[1:4]

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY = Font(name="Arial", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")

def copy_sheet(src, dst):
    for row in src.iter_rows():
        for c in row:
            if c.value is None and not c.has_style:
                continue
            d = dst.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                d.font = copy(c.font)
                d.fill = copy(c.fill)
                d.border = copy(c.border)
                d.alignment = copy(c.alignment)
                d.number_format = c.number_format
    for k, dim in src.column_dimensions.items():
        if dim.width:
            dst.column_dimensions[k].width = dim.width
    for k, dim in src.row_dimensions.items():
        if dim.height:
            dst.row_dimensions[k].height = dim.height
    for m in src.merged_cells.ranges:
        dst.merge_cells(str(m))
    dst.freeze_panes = src.freeze_panes
    if src.auto_filter.ref:
        dst.auto_filter.ref = src.auto_filter.ref
    for dv in src.data_validations.dataValidation:
        ndv = DataValidation(type=dv.type, formula1=dv.formula1, allow_blank=dv.allow_blank)
        ndv.sqref = dv.sqref
        dst.add_data_validation(ndv)
    dst.sheet_view.showGridLines = src.sheet_view.showGridLines

out = openpyxl.load_workbook(DICT)          # dictionary tabs come across as they are
lib = openpyxl.load_workbook(LIB)

LIB_SUMMARY = {
    "Read Me": "How the data library was pulled and what each tab holds",
    "Org Summary": "Edition, licences, users, storage, limits",
    "Objects": "336 objects with data: counts, package, record types, automation per object",
    "Fields": "3,683 fields on 98 profiled objects: type, formula, help text, fill percentage, references",
    "Picklist Values": "6,765 values with defined / in-use / undefined status and record counts",
    "Record Types": "Record types with active flag and business process",
    "Validation Rules": "79 active rules with formula and message",
    "Flows": "279 flow definitions with object, trigger, status, description",
    "Workflows Triggers Approvals": "Classic workflow rules, Apex triggers, approval processes",
    "Scheduled Jobs": "Cron entries with type and schedule",
    "Integrations": "Connected apps, named credentials, remote sites, login clients",
    "Users Access": "Profiles, permission sets, roles, sharing defaults",
    "Data Volumes": "Storage by object",
    "Data Quality": "Fill-rate and orphan-value findings",
    "Reports UI": "Reports, dashboards, list views, Lightning pages",
    "Merge Considerations": "The working list behind the Gaps tab",
    "Field Disposition": "Proposed Keep / Review / Retire per field with reasons and a Decision column",
    "Object Scope": "What moves per object: all, subset, convert, archive, drop",
    "Control Totals": "Reconciliation counts and sums with their queries",
    "Keys Crosswalk": "External identifiers per system with coverage; proposed merge match key",
    "Account Duplicates": "Internal duplicate groups by normalised name and state",
    "Reports Used 365d": "The 855 reports that ran in the last year",
}

# --- merge the library backlog notes into the dictionary backlog -------------
lib_bl = lib["Cleanup Backlog"]
notes = [r[6] for r in lib_bl.iter_rows(min_row=2, values_only=True) if r[0]]
dbl = out["Cleanup Backlog"]
hdr_row = next(r for r in range(1, 12) if dbl.cell(row=r, column=1).value == "#")
ncol = dbl.max_column + 1
h = dbl.cell(row=hdr_row, column=ncol, value="Note")
h.fill, h.font = HDR_FILL, HDR_FONT
for i, note in enumerate(notes, 1):
    c = dbl.cell(row=hdr_row + i, column=ncol, value=note)
    c.font, c.alignment = BODY, WRAP
dbl.column_dimensions[get_column_letter(ncol)].width = 60
dbl.auto_filter.ref = f"A{hdr_row}:{get_column_letter(ncol)}{hdr_row + len(notes)}"

# --- copy the library tabs -----------------------------------------------------
added = []
for name in lib.sheetnames:
    if name == "Cleanup Backlog":
        continue
    new_name = "Library Read Me" if name == "Read Me" else name
    if new_name in out.sheetnames:
        new_name = (new_name + " (lib)")[:31]
    ws = out.create_sheet(new_name)
    copy_sheet(lib[name], ws)
    added.append((new_name, name))

# --- extend the Contents tab -----------------------------------------------------
toc = out["Contents"]
last = toc.max_row
toc.cell(row=last + 2, column=1, value="Data library tabs (field-level and migration detail)").font = Font(name="Arial", bold=True, size=11)
r = last + 3
for new_name, name in added:
    toc.cell(row=r, column=1, value=new_name).font = BODY
    toc.cell(row=r, column=2, value="Data library").font = BODY
    c = toc.cell(row=r, column=3, value=LIB_SUMMARY.get(name, "")); c.font = BODY; c.alignment = WRAP
    r += 1

out.save(OUT)
print("saved", OUT, "tabs:", len(out.sheetnames))
