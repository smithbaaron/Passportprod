#!/usr/bin/env python3
"""Convert the Process & Data Dictionary (markdown export of the Claude Doc)
into a maintainable Excel workbook: one tab per section, companion tabs with
dropdowns, and a change log.

Usage: dictionary_to_xlsx.py <main.md> <process_cards.md> <field_mapping.md> <cleanup_backlog.md> <out.xlsx>
"""
import re, sys, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

MAIN, CARDS, MAPPING, BACKLOG, OUT = sys.argv[1:6]

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
H1 = Font(name="Arial", bold=True, size=16)
H2 = Font(name="Arial", bold=True, size=13, color="1F3864")
BOLD = Font(name="Arial", bold=True, size=10)
BODY = Font(name="Arial", size=10)
QUIET = Font(name="Arial", size=9, italic=True, color="666666")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = openpyxl.Workbook()
wb.remove(wb.active)

def clean(s):
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"`(.+?)`", r"\1", s)
    s = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", s)
    s = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", s)
    return s.strip()

def split_row(line):
    line = line.strip().strip("|")
    return [clean(c) for c in line.split("|")]

def parse(md):
    """Return list of (title, blocks); block = ('p'|'li'|'ol'|'table'|'h3', payload)."""
    sections, cur, blocks = [], None, []
    lines = md.splitlines()
    i = 0
    para = []
    def flush():
        nonlocal para
        if para:
            blocks.append(("p", clean(" ".join(para))))
            para = []
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("# ") and not cur:
            i += 1; continue
        if ln.startswith("## "):
            flush()
            if cur is not None:
                sections.append((cur, blocks))
            cur, blocks = clean(ln[3:]), []
            i += 1; continue
        if ln.startswith("### "):
            flush(); blocks.append(("h3", clean(ln[4:]))); i += 1; continue
        if ln.strip().startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                r = split_row(lines[i])
                if not all(re.fullmatch(r":?-{2,}:?", c.strip()) or c.strip() == "" for c in r):
                    rows.append(r)
                i += 1
            blocks.append(("table", rows)); continue
        m = re.match(r"^\s*[-*]\s+(.*)", ln)
        if m:
            flush(); blocks.append(("li", clean(m.group(1)))); i += 1; continue
        m = re.match(r"^\s*(\d+)\.\s+(.*)", ln)
        if m:
            flush(); blocks.append(("ol", f"{m.group(1)}. {clean(m.group(2))}")); i += 1; continue
        if not ln.strip():
            flush(); i += 1; continue
        para.append(ln.strip()); i += 1
    flush()
    if cur is not None:
        sections.append((cur, blocks))
    return sections

TAB_NAMES = {
    "How to read this document": "Read Me",
    "Passport at a glance": "At a Glance",
    "Core data model": "Data Model",
    "Process 1: Lead to Opportunity": "P1 Lead to Opportunity",
    "Process 2: Bid Intake and Capture Strategy": "P2 Bid Intake & CapStrat",
    "Process 3: Opportunity lifecycle": "P3 Opportunity Lifecycle",
    "Process 4: Contract to Opportunity to Asset": "P4 Contract to Asset",
    "Process 5: Renewals": "P5 Renewals",
    "Process 6: Client Success": "P6 Client Success",
    "Process 7: FOIA request tracking": "P7 FOIA Tracking",
    "One deal traced end to end": "Deal Trace (Racine)",
    "Revenue and finance data": "Revenue Data",
    "Automation and integration inventory": "Automation & Integrations",
    "Key picklists and definitions": "Picklists & Definitions",
    "Migration readiness": "Migration Readiness",
    "Known gaps and merge considerations": "Gaps & Merge Questions",
    "Appendix: the full field dictionary": "Appendix Workbook",
}

def write_section(title, blocks, index):
    name = TAB_NAMES.get(title, re.sub(r"[\[\]:*?/\\]", "", title))[:31]
    ws = wb.create_sheet(name)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 46
    for col in "CDEFGHIJ":
        ws.column_dimensions[col].width = 26
    ws["B1"] = title; ws["B1"].font = H1
    ws["B2"] = f"Section {index} of the Passport Salesforce Process & Data Dictionary. Drawings from the online document are not reproduced here; the text states what they show."
    ws["B2"].font = QUIET; ws["B2"].alignment = WRAP
    ws.merge_cells("B2:J2"); ws.row_dimensions[2].height = 28
    r = 4
    for kind, payload in blocks:
        if kind == "p":
            ws.cell(row=r, column=2, value=payload).font = BODY
            ws.cell(row=r, column=2).alignment = WRAP
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=10)
            ws.row_dimensions[r].height = max(15, 15 * (len(payload) // 165 + 1))
            r += 1
        elif kind in ("li", "ol"):
            ws.cell(row=r, column=1, value="•" if kind == "li" else "").font = BODY
            c = ws.cell(row=r, column=2, value=payload); c.font = BODY; c.alignment = WRAP
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=10)
            ws.row_dimensions[r].height = max(15, 15 * (len(payload) // 165 + 1))
            r += 1
        elif kind == "h3":
            ws.cell(row=r, column=2, value=payload).font = H2; r += 1
        elif kind == "table":
            r += 1
            for ri, row in enumerate(payload):
                for ci, val in enumerate(row):
                    c = ws.cell(row=r, column=2 + ci, value=val)
                    c.border = BORDER; c.alignment = WRAP
                    if ri == 0:
                        c.fill, c.font = HDR_FILL, HDR_FONT
                    else:
                        c.font = BODY
                longest = max((len(v) for v in row), default=0)
                ws.row_dimensions[r].height = max(15, 13 * (longest // 40 + 1))
                r += 1
            r += 1
    ws.sheet_view.showGridLines = False
    return ws

def write_table_tab(name, md, intro_rows, status_col_name=None, status_values=None, widths=None):
    sections = parse(md)
    body = "\n".join(l for l in md.splitlines() if not l.startswith("# "))
    blocks = sections[0][1] if sections else parse("## x\n" + body)[0][1]
    table = next(p for k, p in blocks if k == "table")
    intro = [p for k, p in blocks if k == "p"]
    ws = wb.create_sheet(name)
    ws["A1"] = name; ws["A1"].font = H1
    r = 2
    for p in intro + intro_rows:
        ws.cell(row=r, column=1, value=p).font = QUIET
        ws.cell(row=r, column=1).alignment = WRAP
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(table[0]))
        ws.row_dimensions[r].height = 30; r += 1
    r += 1
    hdr_row = r
    for ri, row in enumerate(table):
        for ci, val in enumerate(row):
            c = ws.cell(row=r, column=1 + ci, value=val)
            c.border = BORDER; c.alignment = WRAP
            if ri == 0:
                c.fill, c.font = HDR_FILL, HDR_FONT
            else:
                c.font = BODY
        r += 1
    last = r - 1
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)
    ws.auto_filter.ref = f"A{hdr_row}:{get_column_letter(len(table[0]))}{last}"
    if status_col_name and status_values:
        col = table[0].index(status_col_name) + 1
        dv = DataValidation(type="list", formula1='"' + ",".join(status_values) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(col)}{hdr_row + 1}:{get_column_letter(col)}{last + 200}")
    ws.sheet_view.showGridLines = False
    return ws

# ------------------------------------------------------------------ contents tab
toc = wb.create_sheet("Contents")
toc["A1"] = "Passport Salesforce: Process & Data Dictionary"; toc["A1"].font = H1
toc["A2"] = f"Maintained workbook version. Source: Claude Doc of the same name; exported {datetime.date.today().isoformat()}. Update the tabs directly and log the change on the Change Log tab."
toc["A2"].font = QUIET; toc["A2"].alignment = WRAP; toc.merge_cells("A2:C2"); toc.row_dimensions[2].height = 30
toc.append([]); toc.append(["Tab", "Section", "What it holds"])
for c in toc[4]:
    c.fill, c.font = HDR_FILL, HDR_FONT
toc.column_dimensions["A"].width = 30; toc.column_dimensions["B"].width = 44; toc.column_dimensions["C"].width = 90

sections = parse(open(MAIN).read())
summaries = {
    "How to read this document": "Audience, scope, conventions, what is left out",
    "Passport at a glance": "Business model, headline counts, teams and the objects they own",
    "Core data model": "The twelve objects and how they link; one Account per customer",
    "Process 1: Lead to Opportunity": "Lead sources, statuses, conversion, rules, gaps",
    "Process 2: Bid Intake and Capture Strategy": "Intake form, bid record, stage path, rules, procurement method, gaps",
    "Process 3: Opportunity lifecycle": "Record types, Types, stages, gates, close rules, fields, automation",
    "Process 4: Contract to Opportunity to Asset": "Ironclad workflow, close, asset per product line, renewal date",
    "Process 5: Renewals": "Auto-created renewals, Renewal Type, stage path, reminders",
    "Process 6: Client Success": "Ownership, cadence, MEDDIC on the account, at-risk journeys, churn",
    "Process 7: FOIA request tracking": "FOIA case record type, statuses, reminders, link to the bid",
    "One deal traced end to end": "City of Racine, WI from account to renewal; the live Suffolk County bid",
    "Revenue and finance data": "NetSuite feed via Celigo, revenue fields, what is not there",
    "Automation and integration inventory": "Flows and workflow rules by object; connected systems",
    "Key picklists and definitions": "Picklist values a mapping needs; business definitions",
    "Migration readiness": "Field and object verdicts, what cannot migrate, identity, reports, ownership",
    "Known gaps and merge considerations": "Data-quality items, structural points, ten questions for Arrive",
    "Appendix: the full field dictionary": "Tabs of the companion data-library workbook",
}
for i, (title, blocks) in enumerate(sections, 1):
    ws = write_section(title, blocks, i)
    toc.append([ws.title, title, summaries.get(title, "")])
for name, what in (("Process Cards", "One row per process with an empty Arrive column"),
                   ("Field Mapping Starter", "Thirty fields to agree on; Arrive columns and a status dropdown"),
                   ("Cleanup Backlog", "Twenty ranked cleanup items with owner, effort and a status dropdown"),
                   ("Change Log", "Who changed what, when")):
    toc.append([name, "", what])
for row in toc.iter_rows(min_row=5):
    for c in row:
        c.font = BODY; c.alignment = WRAP

write_table_tab("Process Cards", open(CARDS).read(), [], widths=[26, 26, 22, 34, 44, 44, 30, 30])
write_table_tab("Field Mapping Starter", open(MAPPING).read(),
                ["Status is a dropdown: Not started, Discussed, Mapped, Gap."],
                status_col_name="Status", status_values=["Not started", "Discussed", "Mapped", "Gap"],
                widths=[20, 30, 16, 52, 24, 30, 14, 30])
write_table_tab("Cleanup Backlog", open(BACKLOG).read(),
                ["Status is a dropdown: Not started, In progress, Blocked, Parked, Done. Add rows at the bottom; keep the numbering."],
                status_col_name="Status", status_values=["Not started", "In progress", "Blocked", "Parked", "Done"],
                widths=[5, 70, 26, 9, 10, 16, 16])

log = wb.create_sheet("Change Log")
log.append(["Date", "Tab / Section", "Change", "By"])
for c in log[1]:
    c.fill, c.font = HDR_FILL, HDR_FONT
for row in (["2026-10-15", "All", "First version written from production metadata and data (read 13 to 15 October 2026)", "Aaron Smith"],
            ["2026-10-16", "Deal Trace, Process Cards, Field Mapping Starter", "Added the Racine walkthrough, process cards and the mapping starter; activity and catalogue figures corrected", "Aaron Smith"],
            ["2026-10-19", "Migration Readiness, Cleanup Backlog", "Added migration readiness findings and the ranked cleanup backlog; companion workbook gained seven migration tabs", "Aaron Smith"],
            [datetime.date.today().isoformat(), "Workbook", "Converted the online document to this maintained workbook, one tab per section", "Aaron Smith"]):
    log.append(row)
for col, w in zip("ABCD", (12, 36, 100, 16)):
    log.column_dimensions[col].width = w
for row in log.iter_rows(min_row=2):
    for c in row:
        c.font = BODY; c.alignment = WRAP
log.freeze_panes = "A2"

wb.save(OUT)
print("saved", OUT, "tabs:", len(wb.sheetnames))
